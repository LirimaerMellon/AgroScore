"""
GET /api/template — скачать Excel-шаблон для заполнения заявок.
"""

from fastapi import APIRouter
from fastapi.responses import StreamingResponse
import io
import openpyxl
import logging

from app.config import TEMPLATE_COLUMNS

logger = logging.getLogger(__name__)
router = APIRouter(prefix="/api", tags=["template"])


@router.get("/template")
def download_template():
    """Генерирует и отдаёт Excel-шаблон с заголовками для inference."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Заявки"

    # Заголовки
    for col_idx, header in enumerate(TEMPLATE_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = openpyxl.styles.Font(bold=True)
        ws.column_dimensions[cell.column_letter].width = max(len(header) + 4, 18)

    # Пример строки (пустая)
    for col_idx in range(1, len(TEMPLATE_COLUMNS) + 1):
        ws.cell(row=2, column=col_idx, value="")

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=agriscore_template.xlsx"},
    )

